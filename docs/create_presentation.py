import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette (Dark Cosmic Theme)
    COLOR_BG = RGBColor(15, 23, 42)         # #0f172a
    COLOR_CARD = RGBColor(30, 41, 59)       # #1e293b
    COLOR_BORDER = RGBColor(51, 65, 85)     # #334155
    COLOR_PLASMA = RGBColor(99, 102, 241)   # #6366f1 (Indigo/Plasma)
    COLOR_AURORA = RGBColor(56, 189, 248)   # #38bdf8 (Cyan/Aurora)
    COLOR_NOVA = RGBColor(236, 72, 153)     # #ec4899 (Pink/Nova)
    COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, subtitle_text):
        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_WHITE
        p.font.name = "Segoe UI"
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(14)
            p2.font.color.rgb = COLOR_AURORA
            p2.font.name = "Segoe UI"

    def add_footer(slide, current_page, total_pages=10):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Vmax.mn | Үл Хөдлөх Хөрөнгийн Платформ   •   Хуудас {current_page} / {total_pages}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.font.name = "Segoe UI"

    def create_card(slide, left, top, width, height, title, items, icon_tag=""):
        # Card Background Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1.5)

        # Content Box
        content_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), height - Inches(0.6))
        tf = content_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Title
        p = tf.paragraphs[0]
        p.text = (icon_tag + " " if icon_tag else "") + title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_AURORA
        p.font.name = "Segoe UI"
        p.space_after = Pt(12)

        # Bullet Items
        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = "• " + item
            p_item.font.size = Pt(13)
            p_item.font.color.rgb = COLOR_TEXT_WHITE
            p_item.font.name = "Segoe UI"
            p_item.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative Center Banner Shape
    banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_CARD
    banner.line.color.rgb = COLOR_PLASMA
    banner.line.width = Pt(2.5)

    tf = banner.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.8)
    tf.margin_top = Inches(0.8)

    p1 = tf.paragraphs[0]
    p1.text = "Vmax.mn"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PLASMA
    p1.font.name = "Segoe UI"
    p1.space_after = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = "Үл Хөдлөх Хөрөнгийн Платформ"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.font.name = "Segoe UI"
    p2.space_after = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "Системийн Иж Бүрэн Давуу Талууд & Шинэлэг Боломжуудын Түлхүүр Танилцуулга"
    p3.font.size = Pt(18)
    p3.font.color.rgb = COLOR_AURORA
    p3.font.name = "Segoe UI"
    p3.space_after = Pt(28)

    p4 = tf.add_paragraph()
    p4.text = "🔑 Сошиал Маркетинг  |  🎬 Видео Зар  |  ⚖️ Зарын Харьцуулалт  |  📊 Аналитик Статистик"
    p4.font.size = Pt(14)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_TEXT_MUTED
    p4.font.name = "Segoe UI"

    # -------------------------------------------------------------
    # SLIDE 2: 🎯 Зорилго & Шинэ Шатанд Гаргах Шийдэл
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "🎯 Платформын Зорилго ба Шинэлэг Шийдэл", "Уламжлалт зарын сайтуудаас ялгарах концепци")
    add_footer(s2, 2)

    create_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9), 
                "Уламжлалт Системүүдийн Асуудал", 
                [
                    "Зарын холбоосыг гараар хуулж фэйсбүүкт нийтлэхэд хүндрэлтэй, цаг их алддаг.",
                    "Сошиалд шэйрлэхэд зарын зураг, үнэ тодорхойгүй харагддаг.",
                    "Зөвхөн текст ба статик зурагтай тул байрны бодит байдлыг мэдрэхэд хэцүү.",
                    "Зарын үзэлтийн дата, сошиал эрэлтийн тоон үзүүлэлт хангалтгүй.",
                    "Заруудыг зэрэгцүүлж үнэ, м.кв-аар харьцуулах боломж байдаггүй."
                ], "❌")

    create_card(s2, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9), 
                "Vmax.mn-ийн Шинэлэг Шийдэл", 
                [
                    "1 товшилтоор сошиалд хуваалцах + Сошиал постер зураг 1 секундэд бэлтгэх.",
                    "Facebook Open Graph Card-аар зарын зураг, үнэ сошиалд тунгалаг харагдана.",
                    "Видео зар ба виртуал аялал хавсаргаж итгэлцлийг 300% нэмэгдүүлнэ.",
                    "Үзэлтийн тоо болон Сошиалд хуваалцсан тоог (Shares) бодит цагаар хянана.",
                    "Заруудыг зэрэгцүүлж 1 м.кв-ын үнэ ба онцлогоор нь шүүн харьцуулна."
                ], "✅")

    # -------------------------------------------------------------
    # SLIDE 3: ⚡ 1. Зарыг Маш Хурдан & Хялбар Оруулах
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "⚡ 1. Зарыг Маш Хурдан & Хялбар Оруулах", "1 минутын дотор зарыг амжилттай нийтлэх ухаалаг Wizard систем")
    add_footer(s3, 3)

    create_card(s3, Inches(0.8), Inches(1.7), Inches(3.64), Inches(4.9),
                "5 Алхамтай Wizard Форм",
                [
                    "Хэрэглэгчийг төөрөгдүүлэхгүй 5 алхамт дараалал.",
                    "1. Төрөл (Зарах/Түрээслэх)",
                    "2. Мэдээлэл (Үнэ, Байршил)",
                    "3. Дэлгэрэнгүй онцлог",
                    "4. Зураг ба Видео",
                    "5. Баталгаажуулалт"
                ], "📝")

    create_card(s3, Inches(4.84), Inches(1.7), Inches(3.64), Inches(4.9),
                "Media Drag & Drop",
                [
                    "Зургуудыг олноор нь нэг дор чирж оруулах боломж.",
                    "Нүүр зургийг хялбархан сонгох.",
                    "Зургийн дарааллыг тааруулах.",
                    "Түргэн ачаалалттай зургийн сан."
                ], "🖼️")

    create_card(s3, Inches(8.88), Inches(1.7), Inches(3.64), Inches(4.9),
                "Улаанбаатарын Бэлэн Бүрхэвч",
                [
                    "Улаанбаатар хотын 9 дүүргийг бэлэн жагсаалтаас сонгоно.",
                    "Хан-Уул, Сүхбаатар, Баянзүрх, Баянгол г.м.",
                    "Байршлын тодорхойлолт автоматаар бөглөгдөнө."
                ], "📍")

    # -------------------------------------------------------------
    # SLIDE 4: 🎬 2. Видео Зар & Виртуал Аялал
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "🎬 2. Видео Зар & Виртуал Аялал (Video Listing)", "Орон сууцны бодит бичлэгийг шууд үзүүлэх боломж")
    add_footer(s4, 4)

    create_card(s4, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9),
                "Олон Төрлийн Видео Хавсаргах",
                [
                    "Шууд видео файл оруулах эсвэл видео холбоос хавсаргах.",
                    "YouTube видео холбоос оруулах.",
                    "TikTok ба Instagram Reels бичлэгүүдийг хавсаргах.",
                    "Зарын нүүр хуудас дээр шууд тоглуулах медиа тоглуулагч."
                ], "🎥")

    create_card(s4, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9),
                "Худалдан Авагчийн Итгэлцэл ба Замын Цаг Хэмнэлт",
                [
                    "Худалдан авагчид заавал биечлэн очих шаардлагагүйгээр виртуал аялал хийнэ.",
                    "Өрөөний зохион байгуулалт, цонхны харагдац, засварыг бодитоор харна.",
                    "Зуучлагч ба эзэмшигчийн дэмий үзлэгт үрэх цагийг 80% хэмнэнэ.",
                    "Борлуулалтын шийдвэр гаргалтыг 3 дахин хурдасгана."
                ], "✨")

    # -------------------------------------------------------------
    # SLIDE 5: ⚖️ 3. Заруудыг Зэрэгцүүлэн Харьцуулах
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "⚖️ 3. Заруудыг Зэрэгцүүлэн Харьцуулах (Property Comparison)", "Шилдэг саналыг сонгох нэгдсэн харьцуулалтын систем")
    add_footer(s5, 5)

    create_card(s5, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9),
                "Зэрэгцүүлэн Шүүх Боломжууд",
                [
                    "2 болон түүнээс дээш зарыг нэг хүснэгтэд зэрэгцүүлэн шүүх.",
                    "Зарын нийт үнийн харьцуулалт.",
                    "1 м.кв-ын нэгж үнийн нарийвчилсан харьцуулалт.",
                    "Талбайн хэмжээ (м.кв) ба өрөөний тооны харьцуулалт.",
                    "Байршил, барилгын давхар, ариун цэврийн өрөөний тоо."
                ], "📊")

    create_card(s5, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9),
                "Худалдан Авагчид Өгөх Үнэ Цэнэ",
                [
                    "Олон таб нээж төөрөх шаардлагагүй болно.",
                    "Хамгийн ашигтай ба боломжийн үнэтэй зарыг шууд тодорхойлно.",
                    "Шийдвэр гаргахад шаардлагатай бүх тоон үзүүлэлт нэг дор.",
                    "Гэр бүл болон хамтран ажиллагсаддаа харьцуулалтыг шэйрлэх."
                ], "💡")

    # -------------------------------------------------------------
    # SLIDE 6: 📢 4. Сошиал Маркетингийн Автоматжуулалт
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "📢 4. Сошиал Маркетингийн Автоматжуулалт", "1 товшилтоор сошиалд хуваалцах & Постер зураг автоматаар бэлтгэх")
    add_footer(s6, 6)

    create_card(s6, Inches(0.8), Inches(1.7), Inches(3.64), Inches(4.9),
                "1-Click Social Sharing",
                [
                    "Facebook timeline & Groups рүү шууд шэйр хийх.",
                    "FB Messenger-ээр хуваалцах.",
                    "Telegram & WhatsApp чат руу илгээх.",
                    "X (Twitter) болон утасны цэсээр шэйрлэх."
                ], "🔗")

    create_card(s6, Inches(4.84), Inches(1.7), Inches(3.64), Inches(4.9),
                "Social Banner Generator",
                [
                    "1200x630px сошиал постер зураг автоматаар бэлтгэнэ.",
                    "Зарын зураг, үнэ, байршил, онцлогийг агуулна.",
                    "1 секундэд PNG файл болгон татан авна.",
                    "FB Post / Story-д бэлэн дизайн."
                ], "🎨")

    create_card(s6, Inches(8.88), Inches(1.7), Inches(3.64), Inches(4.9),
                "Open Graph Protocol",
                [
                    "Фэйсбүүкт зарын холбоос оруулахад зураг нь тод харагдана.",
                    "Зарын нэр ба үнийн мэдээлэл бүрэн карт болно.",
                    "Сошиал орчин дахь даралт (CTR)-ыг 5 дахин өсгөнө."
                ], "🌐")

    # -------------------------------------------------------------
    # SLIDE 7: 📊 5. Зарын Бодит Аналитик ба Статистик
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "📊 5. Зарын Бодит Аналитик ба Статистик", "Зарын эзэн ба зуучлагчдад зориулсан дата хяналт")
    add_footer(s7, 7)

    create_card(s7, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9),
                "Бодит Цагийн Үзүүлэлтүүд",
                [
                    "👁️ Views Counter: Зарыг хэдэн хүн үзсэн бэ?",
                    "🔁 Shares Counter: Зарыг сошиалд хэдэн удаа хуваалцсан бэ?",
                    "Зар бүрийн хуудас дээр болон картан дээр шууд харагдана.",
                    "Зарын эрэлт ба сошиал тархалтыг бодитоор үнэлнэ."
                ], "📈")

    create_card(s7, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9),
                "Seller Dashboard Summary",
                [
                    "Нийт идэвхтэй заруудын тооны нэгдсэн тайлан.",
                    "Нийт үзэлтийн нэгдсэн тоо.",
                    "Нийт сошиал хуваалцсан тооны нэгдсэн статистик.",
                    "Зарыг сунгах, нуух, засах, устгах хурдан удирдлага."
                ], "🖥️")

    # -------------------------------------------------------------
    # SLIDE 8: 🔍 6. Ухаалаг Хайлт ба Дуу Хоолойгоор Хайх
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "🔍 6. Ухаалаг Хайлт ба Дуу Хоолойгоор Хайх", "Voice Search & Advanced Filtering")
    add_footer(s8, 8)

    create_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9),
                "🎙️ Voice Search (Дуу Хоолойгоор Хайх)",
                [
                    "Микрофонд яриад хайлт хийх боломж.",
                    "\"Хан-Уул дүүрэгт 3 өрөө орон сууц\" гэж ярихад автомат шүүлт хийнэ.",
                    "Гар утаснаас хайлт хийхэд маш хурдан бөгөөд хялбар.",
                    "Монгол хэлний яриаг таних технологи."
                ], "🎙️")

    create_card(s8, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9),
                "🎛️ Нарийвчилсан Шүүлтүүрүүд",
                [
                    "Зарын төрөл: Зарах / Түрээслэх.",
                    "Ангилал: Орон сууц, Хаус, Газар, Оффис / Коммершиал.",
                    "Үнийн хязгаар (Доод үнэ - Дээд үнэ).",
                    "Талбайн хэмжээ (Доод м.кв - Дээд м.кв).",
                    "Улаанбаатар хотын 9 дүүргийн шүүлтүүр."
                ], "🎛️")

    # -------------------------------------------------------------
    # SLIDE 9: ❤️ 7. Хадгалсан Зарууд & Сүүлд Үзсэн Түүх
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "❤️ 7. Хадгалсан Зарууд & Сүүлд Үзсэн Түүх", "Favorites & Recent Visits History")
    add_footer(s9, 9)

    create_card(s9, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9),
                "❤️ Хадгалсан Зарууд (Favorites)",
                [
                    "Зүрхэн товчоор сонирхсон заруудаа 1 товшилтоор хадгалах.",
                    "Профайл хэсгийн \"Хадгалсан зарууд\" таб-аас эргэн үзэх.",
                    "Үнийн өөрчлөлт болон зарын статусыг хянах.",
                    "Бүх төхөөрөмж дээр хадгалагдсан хэвээр байна."
                ], "❤️")

    create_card(s9, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9),
                "🕒 Сүүлд Үзсэн Түүх (Recent Visits)",
                [
                    "Үзсэн заруудын түүх автоматаар хадгалагдана.",
                    "Зарын хуудас бүрийн доор карусель байдлаар харагдана.",
                    "Өмнө нь үзэж байсан зараа хайхгүйгээр шууд эргэн шилжих.",
                    "Хэрэглэгчийн цагийг хэмнэсэн ухаалаг UI санамж."
                ], "🕒")

    # -------------------------------------------------------------
    # SLIDE 10: 🌐 8. Олон Хэл ба Системийн Өрсөлдөх Давуу Тал
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "🌐 8. Олон Хэл ба Системийн Өрсөлдөх Давуу Тал", "Яагаад Vmax.mn-ийг сонгох ёстой вэ?")
    add_footer(s10, 10)

    create_card(s10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.9),
                "Олон Хэл & Технологийн Хурд",
                [
                    "Монгол ба Англи (MN / EN) хэлний сонголттай.",
                    "Гадаадын иргэд ба хөрөнгө оруулагчдад нээлттэй.",
                    "Single Page Application (SPA) - маш хурдан ажиллагаа.",
                    "Dark Cosmic өнгө төрх - Орчин үеийн премиум дизайн."
                ], "🌐")

    create_card(s10, Inches(6.933), Inches(1.7), Inches(5.6), Inches(4.9),
                "Товч Дүгнэлт ба Давуу Тал",
                [
                    "✅ 1 минутад зар нийтлэх & 1 секундэд сошиал постер бүтээх.",
                    "✅ Видео зар хавсаргаж худалдан авагчийн итгэлцлийг олох.",
                    "✅ Заруудыг зэрэгцүүлж 1 м.кв-ын үнээр нь харьцуулах.",
                    "✅ Үзэлт ба шэйрийн датагаар маркетингаа удирдах."
                ], "🏆")

    # Save to file
    output_path = os.path.join(os.getcwd(), "Vmax_mn_Platform_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
